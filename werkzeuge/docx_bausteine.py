"""Bausteine für Word-Materialien (Arbeitsblätter, Stundenentwürfe, Klassenarbeiten).

Benutzung in einem Erzeugungs-Skript:

    import sys; sys.path.insert(0, "werkzeuge")
    from docx_bausteine import *
    d = neues_dokument()
    titel(d, "Quadratwurzeln", "Arbeitsblatt 1")
    aufgabe(d, 1, "Berechne.", sterne=1, teile=["√49", "√1,21"], spalten=2, platz=2)
    merkkasten(d, "Merke", ["Die Quadratwurzel √a ist ..."])
    d.save("reihen/01-wurzeln/ab1.docx")

Alles ist schwarz-weiß-tauglich (Graustufen). Farbe nur über farbe=... einschalten.
Markup in Texten: **fett**, _kursiv_.
"""
import math
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GRAU = "E7E6E6"
SCHRIFT = "Calibri"


# ---------------------------------------------------------------- Grundlagen

def neues_dokument(quer=False, rand_cm=1.8, schriftgroesse=12):
    d = Document()
    s = d.sections[0]
    if quer:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
    for r in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, r, Cm(rand_cm))
    st = d.styles["Normal"]
    st.font.name = SCHRIFT
    st.font.size = Pt(schriftgroesse)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SCHRIFT)
    st.paragraph_format.space_after = Pt(4)
    for h in ("Heading 1", "Heading 2", "Heading 3"):
        d.styles[h].font.name = SCHRIFT
        d.styles[h].font.color.rgb = RGBColor(0, 0, 0)
    return d


def _text(p, text, groesse=None, fett=False):
    """Fügt Text mit **fett** / _kursiv_ Markup an Absatz p an."""
    for teil in re.split(r"(\*\*[^*]+\*\*|_[^_]+_)", text):
        if not teil:
            continue
        if teil.startswith("**"):
            r = p.add_run(teil[2:-2]); r.bold = True
        elif teil.startswith("_") and teil.endswith("_") and len(teil) > 2:
            r = p.add_run(teil[1:-1]); r.italic = True
        else:
            r = p.add_run(teil)
        if fett:
            r.bold = True
        if groesse:
            r.font.size = Pt(groesse)
    return p


def absatz(d_oder_zelle, text="", groesse=None, fett=False, zentriert=False, abstand_nach=4):
    p = d_oder_zelle.add_paragraph()
    _text(p, text, groesse, fett)
    p.paragraph_format.space_after = Pt(abstand_nach)
    if zentriert:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def liste(d_oder_zelle, punkte, nummeriert=False):
    for t in punkte:
        p = d_oder_zelle.add_paragraph(style="List Number" if nummeriert else "List Bullet")
        _text(p, t)


def ueberschrift(d, text, ebene=2):
    return d.add_heading(text, level=ebene)


def titel(d, haupt, unter=None):
    p = d.add_paragraph()
    r = p.add_run(haupt); r.bold = True; r.font.size = Pt(18)
    p.paragraph_format.space_after = Pt(2)
    if unter:
        p2 = d.add_paragraph(); r2 = p2.add_run(unter); r2.font.size = Pt(12); r2.italic = True
    _linie(p if not unter else p2)


def seitenumbruch(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _linie(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "4"), ("color", "000000")):
        bot.set(qn("w:" + k), v)
    b.append(bot); pPr.append(b)


def _schattierung(zelle, farbe):
    tcPr = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), farbe)
    tcPr.append(shd)


def _rahmen(tabelle, art="single", staerke=8):
    tbl = tabelle._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for kante in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + kante)
        e.set(qn("w:val"), art); e.set(qn("w:sz"), str(staerke)); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def _breiten(tabelle, breiten_cm):
    tabelle.autofit = False
    for i, w in enumerate(breiten_cm):
        tabelle.columns[i].width = Cm(w)
    for row in tabelle.rows:
        for i, w in enumerate(breiten_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


# ---------------------------------------------------------------- Bausteine

def sterne_text(n):
    return "★" * n + "☆" * (3 - n) if n else ""


def aufgabe(d, nr, text, sterne=0, teile=None, spalten=1, platz=0, punkte=None, afb=None):
    """Aufgabe mit Nummer, Niveau-Sternen, optionalen Teilaufgaben a), b) ...

    platz: Anzahl Schreibzeilen (Kästchen-Ersatz) nach der Aufgabe.
    punkte/afb: für Klassenarbeiten (erscheint rechts bzw. nur im Erwartungshorizont).
    """
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True
    kopf = f"Aufgabe {nr}"
    if sterne:
        kopf += f"  {sterne_text(sterne)}"
    r = p.add_run(kopf); r.bold = True
    if punkte is not None:
        r2 = p.add_run(f"\t({_zahl(punkte)} P.)"); r2.italic = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.3), alignment=2)
    p2 = d.add_paragraph(); _text(p2, text)
    p2.paragraph_format.keep_with_next = bool(teile)
    if teile:
        buchst = "abcdefghijklmnopqrstuvwxyz"
        if spalten > 1:
            zeilen = math.ceil(len(teile) / spalten)
            t = d.add_table(rows=zeilen, cols=spalten)
            for k, teil in enumerate(teile):
                z = t.cell(k % zeilen, k // zeilen)
                _text(z.paragraphs[0], f"{buchst[k]})  {teil}")
        else:
            for k, teil in enumerate(teile):
                q = d.add_paragraph(); q.paragraph_format.left_indent = Cm(0.6)
                _text(q, f"{buchst[k]})  {teil}")
    schreibplatz(d, platz)


def schreibplatz(d, zeilen):
    """Punktierte Schreiblinien über die volle Breite."""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    for _ in range(zeilen):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(24)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run("\t"); r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def merkkasten(d, ueberschrift_text, zeilen, farbe=None):
    """Grau hinterlegter, umrandeter Kasten (Merke / Wissen / Beispiel)."""
    t = d.add_table(rows=1, cols=1); _rahmen(t, staerke=12)
    z = t.cell(0, 0); _schattierung(z, farbe or GRAU)
    _text(z.paragraphs[0], ueberschrift_text, fett=True)
    for zeile in zeilen:
        absatz(z, zeile, abstand_nach=2)
    d.add_paragraph()
    return t


def tabelle(d, kopf, zeilen, breiten_cm=None, kopf_grau=True, schrift=None):
    t = d.add_table(rows=1 + len(zeilen), cols=len(kopf)); _rahmen(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, k in enumerate(kopf):
        c = t.cell(0, i); _text(c.paragraphs[0], k, schrift, fett=True)
        if kopf_grau:
            _schattierung(c, GRAU)
    for r, zeile in enumerate(zeilen, start=1):
        for i, wert in enumerate(zeile):
            c = t.cell(r, i)
            teile = str(wert).split("\n")
            _text(c.paragraphs[0], teile[0], schrift)
            for rest in teile[1:]:
                absatz(c, rest, schrift, abstand_nach=0)
    if breiten_cm:
        _breiten(t, breiten_cm)
    d.add_paragraph()
    return t


def tippkarten(d, karten, titel_text="Tippkarten", spalten=2):
    """karten: Liste von (Überschrift, Text). Gestrichelte Schnittränder zum Ausschneiden."""
    absatz(d, f"✂  {titel_text}", fett=True)
    zeilen = math.ceil(len(karten) / spalten)
    t = d.add_table(rows=zeilen, cols=spalten); _rahmen(t, art="dashed", staerke=6)
    for k, (kopf, text) in enumerate(karten):
        z = t.cell(k // spalten, k % spalten)
        _text(z.paragraphs[0], kopf, fett=True)
        for zeile in text.split("\n"):
            absatz(z, zeile, abstand_nach=2)
        absatz(z, "")
    d.add_paragraph()


def verlaufsplan(d, zeilen):
    """zeilen: Liste von (Phase, Zeit, Unterrichtsgeschehen, Sozialform/Methode, Medien)."""
    return tabelle(d, ["Phase", "Zeit", "Unterrichtsgeschehen", "Sozialform / Methode", "Medien"],
                   zeilen, breiten_cm=[2.6, 1.6, 8.0, 3.0, 2.6], schrift=10)


# ---------------------------------------------------------------- Bewertung

NOTEN_KA = [(90, "sehr gut (1)"), (75, "gut (2)"), (60, "befriedigend (3)"),
            (45, "ausreichend (4)"), (20, "mangelhaft (5)"), (0, "ungenügend (6)")]


def _zahl(x):
    return (f"{x:.1f}".rstrip("0").rstrip(".")).replace(".", ",")


def punktgrenzen(gesamt):
    """Mindestpunktzahl je Note, aufgerundet auf halbe Punkte."""
    return [(note, math.ceil(prozent / 100 * gesamt * 2 - 1e-9) / 2) for prozent, note in NOTEN_KA]


def notentabelle(d, gesamt):
    grenzen = punktgrenzen(gesamt)
    zeilen = []
    for i, (note, ab) in enumerate(grenzen):
        bis = gesamt if i == 0 else grenzen[i - 1][1] - 0.5
        zeilen.append([note, f"{_zahl(ab)} – {_zahl(bis)}", f"ab {NOTEN_KA[i][0]} %"])
    return tabelle(d, ["Note", "Punkte", "Prozent"], zeilen, breiten_cm=[5, 4, 3])
