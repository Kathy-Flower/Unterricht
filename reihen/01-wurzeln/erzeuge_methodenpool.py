"""Methodenpool Reelle Zahlen: Domino, Sortierspiel, Partnercheck, Fehlerdetektiv,
Gruppenpuzzle (Expertenblätter A–D = Pflichtstationen der Lerntheke), Lerntheke-Laufzettel.

Aufruf aus dem Repo-Wurzelverzeichnis:  python3 reihen/01-wurzeln/erzeuge_methodenpool.py
"""
import random
import sys
sys.path.insert(0, "werkzeuge")
from docx_bausteine import *
from docx_bausteine import _rahmen, _schattierung, _text, _nicht_trennen

d = neues_dokument(schriftgroesse=11)


def steckbrief(zeilen):
    t = d.add_table(rows=len(zeilen), cols=2); _rahmen(t)
    for i, (k, v) in enumerate(zeilen):
        _schattierung(t.cell(i, 0), GRAU)
        _text(t.cell(i, 0).paragraphs[0], k, fett=True)
        _text(t.cell(i, 1).paragraphs[0], v)
    from docx.shared import Cm
    for r in t.rows:
        r.cells[0].width, r.cells[1].width = Cm(3.6), Cm(13.8)
    d.add_paragraph()


def karten(inhalte, spalten=3, gross=False):
    """Schneidekarten: inhalte = Liste von Listen von Zeilen (erste Zeile fett)."""
    zeilen = math.ceil(len(inhalte) / spalten)
    t = d.add_table(rows=zeilen, cols=spalten); _rahmen(t, art="dashed", staerke=6)
    for r in t.rows:
        _nicht_trennen(r)
    for k, zl in enumerate(inhalte):
        z = t.cell(k // spalten, k % spalten)
        absatz(z, zl[0], groesse=20 if gross else None, fett=True, zentriert=True, abstand_nach=2)
        for zeile in zl[1:]:
            absatz(z, zeile, zentriert=True, abstand_nach=2)
        absatz(z, "", abstand_nach=4)
    d.add_paragraph()


def neues_blatt(haupt, unter=None):
    titel(d, haupt, unter)
    d.paragraphs[-2 if unter else -1].paragraph_format.page_break_before = True


def satzbausteine(zeilen, kopf="Wortspeicher und Satzbausteine"):
    merkkasten(d, kopf, zeilen)


# ===================================================================== Lehrkraftteil
titel(d, "Methodenpool: Reelle Zahlen (Wurzeln)", "Klasse 9 – Arbeitsaufträge mit Methode, fertig zum Kopieren")

absatz(d, "Der Pool ergänzt deine Arbeitsblätter (Einstiege, Trainings, Tandem, Partnerspiel „Wurzelpaar“). "
          "Alle Aufgaben sind neu, nachgerechnet und ohne Buchverweise. Jede Methode hat einen Steckbrief, "
          "eine Einstiegsfrage mit Lebensweltbezug und eine Lösungsseite. "
          "**Die Expertenblätter A–D (Gruppenpuzzle) sind zugleich die Pflichtstationen der Lerntheke** – "
          "du kopierst sie nur einmal.")

ueberschrift(d, "Überblick", 3)
tabelle(d, ["Methode", "Inhalt", "Dauer", "Sozialform", "Vorschlag für den Einsatz"], [
    ["**1 Wurzel-Domino** (Kartenspiel)", "Quadratwurzeln, x² = a, einfache Wurzelgesetze", "10–15 min", "PA / 3er",
     "Einstieg Fragestunde Mi 04.11. oder Lerntheke"],
    ["**2 Wo wohnt die Zahl?** (Sortier-Kartenspiel)", "Zahlbereiche ℕ ⊂ ℤ ⊂ ℚ ⊂ ℝ", "15–20 min", "4er",
     "Sicherung nach Stunde 12 (Do 08.10.) oder Lerntheke"],
    ["**3 Partnercheck**", "Wurzelgesetze, teilweises Wurzelziehen, binomische Formeln", "15–20 min", "PA",
     "Mi 14.10. (Brücke Terme) oder Mi 04.11."],
    ["**4 Fehlerdetektiv** „Klassenchat“", "typische Fehler aus allen Teilthemen", "10–15 min", "EA → PA",
     "Einstieg Lerntheke oder Fragestunde"],
    ["**5 Gruppenpuzzle** „Wurzel-Expert*innen“", "Wiederholung aller vier Teilthemen", "90 min", "Experten- / Stammgruppen",
     "Alternative für Do 15.10."],
    ["**6 Lerntheke** mit Laufzettel", "Pflicht: Expertenblätter A–D, Wahl: Methoden 1–4 + dein Partnerspiel",
     "90 min", "EA / PA / GA", "**Empfehlung für Do 15.10.** (Vertiefung & Diagnose vor der KA)"],
], breiten_cm=[3.9, 4.6, 1.7, 2.4, 4.8], schrift=10)

absatz(d, "**Differenzierung** überall: ★ Mindeststandard · ★★ Standard · ★★★ Experte. "
          "Lösungen liegen jeweils auf einer eigenen Seite (Lösungsstation oder Rückseite). "
          "**Gruppen für 31 SuS**: 7 Vierer + 1 Dreier; beim Domino 14 Paare + 1 Dreier.")

# ---------------------------------------------------------------- Steckbriefe
ueberschrift(d, "1  Wurzel-Domino (Kartenspiel)", 3)
steckbrief([
    ("Einstiegsfrage", "„Ein quadratischer Instagram-Post hat 1 166 400 Pixel. Wie breit ist er – ohne Taschenrechner?“ "
                       "Die Karte liegt im Spiel; wer sie zieht, muss eine Strategie finden (1 166 400 = 11 664 · 100)."),
    ("Ablauf", "Karten mischen und verteilen. Eine beliebige Karte beginnt. Wer die **Aufgabe** der letzten Karte "
               "löst, legt die Karte mit dem passenden **Ergebnis** an. Die Kette schließt sich zum Ring – "
               "passt die letzte Aufgabe zur ersten Karte, ist alles richtig (Selbstkontrolle)."),
    ("Material", "1 Kartensatz (16 Karten) pro Paar, ohne Taschenrechner."),
    ("Kompetenzen", "(7), (9), Ope-1, Ope-4"),
    ("Differenzierung", "Schnelle Paare: Ring rückwärts erklären oder selbst zwei neue Karten erfinden, die in den Ring passen."),
])

ueberschrift(d, "2  Wo wohnt die Zahl? (Sortier-Kartenspiel)", 3)
steckbrief([
    ("Einstiegsfrage", "„Der Taschenrechner zeigt für √2 die Zahl 1,4142… – ist √2 dann ein Bruch?“ "
                       "Die Karte „1,4142 (Taschenrechner-Anzeige)“ ist bewusst als Streitkarte im Spiel."),
    ("Ablauf", "Reihum Karte ziehen, in das **kleinste passende** Feld des Spielplans legen und mit einem "
               "Satzbaustein begründen (1 Punkt). Alle anderen dürfen „**Einspruch!**“ rufen. Die Lösungskarte "
               "entscheidet: berechtigter Einspruch = Punkt für die Einsprechenden, Karte wird umgelegt."),
    ("Material", "Spielplan (A4, besser auf A3 kopiert), 20 Zahlkarten, gefaltete Lösungskarte, je 4er-Gruppe."),
    ("Kompetenzen", "(2), Arg-2, Arg-4, Kom-3, Kom-6"),
    ("Differenzierung", "Wortspeicher auf dem Spielplan; ★★★: zu jedem Feld eine eigene „gemeine“ Karte erfinden."),
])

ueberschrift(d, "3  Partnercheck", 3)
steckbrief([
    ("Einstiegsfrage", "„Ihr spielt verschiedene Level, aber am Ende muss derselbe Score herauskommen. "
                       "Wenn nicht: Wer hat sich verrechnet?“"),
    ("Ablauf", "A rechnet Spalte A, B rechnet Spalte B – **in jeder Zeile kommt dasselbe Ergebnis heraus**. "
               "Nach jeder Zeile vergleichen. Unterschiedlich? Gemeinsam den Fehler suchen, erst dann weiter."),
    ("Material", "1 Blatt pro Paar (in der Mitte falzen), Lösungen an der Lösungsstation."),
    ("Kompetenzen", "(7), Ope-5, Arg-9, Kom-8"),
    ("Differenzierung", "Zeilen 1–5 ★, 6–8 ★★, 9–10 ★★★ (binomische Formeln, Variable). Mindestziel: Zeile 6."),
])

ueberschrift(d, "4  Fehlerdetektiv „Klassenchat“", 3)
steckbrief([
    ("Einstiegsfrage", "„Am Abend vor der Klassenarbeit werden im Klassenchat Lösungen gepostet. Kann man denen trauen?“"),
    ("Ablauf", "Einzeln: Fehler markieren und richtige Lösung notieren (5 min). Paar: Fehler mit Satzbaustein "
               "erklären (5 min). Plenum: „Welcher Fehler ist der gefährlichste für die Klassenarbeit?“"),
    ("Material", "1 Blatt pro Person."),
    ("Kompetenzen", "Arg-9, Arg-10, Kom-6"),
    ("Differenzierung", "Nachrichten 1–4 ★, 5–7 ★★, 8 ★★★; Satzbausteine auf dem Blatt."),
])

ueberschrift(d, "5  Gruppenpuzzle „Wurzel-Expert*innen“ (90 min)", 3)
steckbrief([
    ("Leitfrage", "„Wir wollen ein quadratisches Banner mit 2 m² Fläche drucken. Wie lang ist die Seite – exakt, "
                  "auf cm genau, und ist das ein Bruch?“ Die SuS merken: Dafür braucht man alle vier Teilthemen → Gruppenpuzzle."),
    ("Gruppen", "7 Stammgruppen à 4 + 1 à 3. Jedes Mitglied wird Expert*in für A, B, C oder D. Pro Thema zwei "
                "Expertengruppen (à 3–4). **Steuern**: A (Quadratwurzeln) ist der leichteste Einstieg – "
                "leistungsschwächere SuS bewusst dort einsetzen. Die Dreiergruppe liest Thema D gemeinsam "
                "(Infokasten + Kontrollfrage)."),
    ("Material", "Expertenblätter A–D (je 8 Kopien), Abschluss-Challenge (8 Kopien), Lösungen."),
    ("Kompetenzen", "(2), (6), (7), (9), Kom-4, Kom-8, Arg-5"),
])
verlaufsplan(d, [
    ["Einstieg", "10'", "Banner-Problem (Beamer). SuS sammeln, was sie wissen müssen (exakter Wert, Näherung, "
                        "Bruch?, Rechnen mit Wurzeln) → Leitfrage an der Tafel. Gruppeneinteilung.", "Plenum", "Beamer"],
    ["Expertenphase", "25'", "Infokasten lesen, ★-Aufgaben einzeln (10'), dann ★★/★★★ und Erklärzettel in der "
                             "Expertengruppe (15'). Lösungen an der Lösungsstation.", "EA → GA", "Expertenblätter A–D"],
    ["Stammgruppe", "35'", "Jede*r erklärt sein Thema (je ca. 6') und stellt die Kontrollfrage. Danach "
                           "Abschluss-Challenge gemeinsam.", "GA", "Erklärzettel, Challenge"],
    ["Sicherung & Reflexion", "15'", "Challenge im Plenum (eine Gruppe präsentiert, andere ergänzen). Beantwortung "
                                    "der Leitfrage. Ampel: „Welches Thema übe ich vor der KA noch?“", "Plenum", "Dokumentenkamera"],
    ["Puffer", "5'", "Raumwechsel / Umsetzen", "", ""],
])

ueberschrift(d, "6  Lerntheke mit Laufzettel (90 min)", 3)
steckbrief([
    ("Einstiegsfrage", "Klassenchat-Nachricht „√(16 + 9) = 4 + 3 = 7“ an den Beamer: „Stimmt das? Und welche Fehler "
                       "dürfen mir in der Klassenarbeit nicht passieren?“ → SuS schätzen sich auf dem Laufzettel ein."),
    ("Stationen", "**Pflicht** P1–P4 = Expertenblätter A–D (jeweils mindestens die ★-Aufgaben). "
                  "**Wahl** W1 Domino · W2 Sortierspiel · W3 Partnercheck · W4 Fehlerdetektiv · "
                  "W5 dein Partnerspiel „Wer findet schneller ein Wurzelpaar?“."),
    ("Organisation", "Pflichtstationen als Stapel am Pult (je 8 Kopien), Wahlstationen auf Fensterbänken, "
                     "Lösungsstation am Pult. Wer bei der Selbsteinschätzung „unsicher“ ankreuzt, beginnt mit "
                     "der zugehörigen Pflichtstation."),
    ("Kompetenzen", "(2), (6), (7), (9), Selbstdiagnose"),
])
verlaufsplan(d, [
    ["Einstieg", "10'", "Klassenchat-Nachricht am Beamer, kurzer Austausch. Leitfrage: „Was muss ich bis zur KA "
                        "noch üben?“ Selbsteinschätzung (vorher) auf dem Laufzettel.", "Plenum → EA", "Beamer, Laufzettel"],
    ["Arbeitsphase", "65'", "Pflicht- und Wahlstationen in eigenem Tempo, Selbstkontrolle an der Lösungsstation. "
                            "Lehrkraft: gezielte Förderung, v. a. bei P1/P4.", "EA / PA / GA", "Stationen, Lösungen"],
    ["Sicherung & Reflexion", "15'", "Selbsteinschätzung (nachher), Exit-Ticket auf dem Laufzettel. "
                                    "Hinweis auf Übungsmöglichkeiten in den Ferien.", "EA → Plenum", "Laufzettel"],
])

# ===================================================================== 1 Domino
DOMINO = [  # (Aufgabe, Ergebnis) in Ringreihenfolge
    ("√144", "12"), ("√0,09", "0,3"), ("√(49/64)", "7/8"), ("√2 · √8", "4"), ("√50 : √2", "5"),
    ("Seitenlänge (in m) eines quadratischen Zimmers mit 2,25 m²", "1,5"),
    ("√(6² + 8²)", "10"), ("positive Lösung von x² = 169", "13"), ("(√7)²", "7"), ("√1,21", "1,1"),
    ("√0,0004", "0,02"), ("√12 · √3", "6"),
    ("Breite (in Pixel) eines quadratischen Instagram-Posts mit 1 166 400 Pixeln", "1080"),
    ("√(16 · 25)", "20"), ("Breite (in Pixel) eines quadratischen Pixel-Art-Sprites aus 256 Pixeln", "16"),
    ("√72 (teilweise Wurzel ziehen)", "6√2"),
]
n = len(DOMINO)
karten_dom = [[DOMINO[(i - 1) % n][1], "─────────", "Aufgabe: " + DOMINO[i][0]] for i in range(n)]
random.Random(9).shuffle(karten_dom)
neues_blatt("Wurzel-Domino", "Kartenspiel zu zweit oder zu dritt – ohne Taschenrechner")
absatz(d, "**So geht’s:** Mischt die Karten und verteilt sie. Legt eine beliebige Karte in die Mitte. "
          "Löst die **Aufgabe** auf dieser Karte. Wer die Karte mit dem passenden **Ergebnis** (oben) hat, legt sie an "
          "und erklärt kurz den Rechenweg. Am Ende muss sich die Kette zu einem **Ring** schließen: "
          "Die Aufgabe der letzten Karte führt zum Ergebnis der ersten Karte.")
absatz(d, "★★★ **Extra:** Erfindet zwei neue Karten, die ihr zwischen zwei vorhandene Karten in den Ring einbauen könnt.")
karten(karten_dom, spalten=4, gross=True)

neues_blatt("Lösung: Wurzel-Domino", "Ringreihenfolge (Start beliebig)")
tabelle(d, ["Nr.", "Aufgabe", "Ergebnis", "Rechenweg / Tipp"], [
    [str(i + 1), a, e, w] for i, ((a, e), w) in enumerate(zip(DOMINO, [
        "12² = 144", "0,3² = 0,09", "√49/√64", "√16", "√25", "1,5² = 2,25", "√100", "13² = 169 (−13 ist die zweite Lösung)",
        "Wurzel und Quadrat heben sich auf", "1,1² = 1,21", "0,02² = 0,0004", "√36",
        "√11 664 · √100 = 108 · 10", "√16 · √25 = 4 · 5", "16² = 256", "√36 · √2"]))
], breiten_cm=[1, 7.5, 2, 6.5], schrift=10)

# ===================================================================== 2 Sortierspiel
SORT = [  # (Karte, Feld, Begründung)
    ("√49", "ℕ", "√49 = 7"), ("12/4", "ℕ", "12/4 = 3"), ("√2 · √8", "ℕ", "= √16 = 4"), ("√12 : √3", "ℕ", "= √4 = 2"),
    ("(−3)²", "ℕ", "= 9"), ("−√16", "ℤ", "= −4"), ("−7", "ℤ", "negative ganze Zahl"), ("1 − √9", "ℤ", "= 1 − 3 = −2"),
    ("0,25", "ℚ", "= 1/4"), ("−3/5", "ℚ", "Bruch"), ("0,3̅", "ℚ", "periodisch, = 1/3"), ("√0,81", "ℚ", "= 0,9 = 9/10"),
    ("√(4/9)", "ℚ", "= 2/3"), ("1,4142\n(Taschenrechner-Anzeige)", "ℚ", "abbrechend: 14 142/10 000 – aber nicht genau √2!"),
    ("√2", "ℝ \\ ℚ", "2 ist keine Quadratzahl"), ("π", "ℝ \\ ℚ", "unendlich, nicht periodisch"),
    ("√20", "ℝ \\ ℚ", "20 ist keine Quadratzahl"), ("0,1010010001…", "ℝ \\ ℚ", "Muster, aber nicht periodisch"),
    ("1 + √3", "ℝ \\ ℚ", "rational + irrational ist irrational"), ("√(−9)", "keine reelle Zahl", "keine Zahl quadriert ergibt −9"),
]
neues_blatt("Wo wohnt die Zahl?", "Sortier-Kartenspiel für 4 Personen – Spielplan")
absatz(d, "**So geht’s:** Reihum zieht ihr eine Karte und legt sie in das **kleinste** Feld, in das die Zahl passt. "
          "Begründe mit einem Satzbaustein – dann gibt es 1 Punkt. Wer anderer Meinung ist, ruft „**Einspruch!**“. "
          "Die Lösungskarte entscheidet: Ist der Einspruch berechtigt, bekommt ihn die einsprechende Person als Punkt "
          "und die Karte wird umgelegt.")

# Spielplan: verschachtelte Tabellen ℝ ⊃ ℚ ⊃ ℤ ⊃ ℕ
from docx.shared import Cm, Pt
plan = d.add_table(rows=1, cols=2); _rahmen(plan, staerke=12)
aussen = plan.cell(0, 0)
_text(aussen.paragraphs[0], "ℝ  reelle Zahlen  (hier hinein: irrationale Zahlen)", fett=True)
for _ in range(2):
    absatz(aussen, "")
zelle = aussen
for name, grau, zeilen, breite in [("ℚ  rationale Zahlen", "F2F2F2", 4, 12.8), ("ℤ  ganze Zahlen", "E7E6E6", 4, 12.0),
                                   ("ℕ  natürliche Zahlen", "D0CECE", 5, 11.2)]:
    t = zelle.add_table(rows=1, cols=1); _rahmen(t, staerke=10)
    t.autofit = False
    zelle = t.cell(0, 0); _schattierung(zelle, grau)
    t.columns[0].width = zelle.width = Cm(breite)
    _text(zelle.paragraphs[0], name, fett=True)
    for _ in range(zeilen):
        absatz(zelle, "")
    if name.startswith("ℕ"):
        break
for c in (aussen,):
    absatz(c, "")
rechts = plan.cell(0, 1)
_text(rechts.paragraphs[0], "keine reelle Zahl", fett=True)
plan.autofit = False
plan.columns[0].width = aussen.width = Cm(13.6)
plan.columns[1].width = rechts.width = Cm(3.8)
d.add_paragraph()
satzbausteine([
    "**Wortspeicher:** natürliche / ganze / rationale / irrationale Zahl · Bruch · abbrechend · periodisch · Quadratzahl · liegt in",
    "„… liegt in ℕ, weil … = … eine natürliche Zahl ist.“",
    "„… ist rational, weil man die Zahl als Bruch … schreiben kann.“",
    "„… ist irrational, weil … keine Quadratzahl ist.“ / „… weil die Dezimalzahl weder abbricht noch periodisch ist.“",
    "„… ist keine reelle Zahl, weil keine Zahl quadriert … ergibt.“",
])

neues_blatt("Wo wohnt die Zahl? – Zahlkarten", "ausschneiden")
karten([k.split("\n") for k, _, _ in SORT], spalten=4, gross=True)
neues_blatt("Wo wohnt die Zahl? – Lösungskarte", "für die Spielleitung oder gefaltet in die Mitte legen")
absatz(d, "✂  **Lösungskarte** (an der gestrichelten Linie falten, erst bei „Einspruch!“ aufklappen)", fett=False)
tabelle(d, ["Karte", "Feld", "Begründung"], [[k.replace("\n", " "), f, b] for k, f, b in SORT],
        breiten_cm=[5.2, 3.4, 8.8], schrift=9)

# ===================================================================== 3 Partnercheck
PC = [
    ("√3 · √12", "√2 · √18", "6"),
    ("√75 : √3", "√200 : √8", "5"),
    ("Seitenlänge eines quadratischen Beets mit 50 m² (exakt)", "3√2 + 2√2", "5√2 (≈ 7,07)"),
    ("√27 (teilweise Wurzel ziehen)", "√12 + √3", "3√3"),
    ("√0,16 · √25", "√0,5 · √8", "2"),
    ("(√5)² + (√3)²", "√(2 · 32)", "8"),
    ("√48 − √12", "√12 (teilweise Wurzel ziehen)", "2√3"),
    ("(√5 − 1)(√5 + 1)", "(√7 − √3)(√7 + √3)", "4"),
    ("(1 + √2)²", "√2 · (√2 + 2) + 1", "3 + 2√2"),
    ("√(18x²)   (x ≥ 0)", "3 · √(2x²)   (x ≥ 0)", "3x√2"),
]
PC_WEG = [
    ("√36 = 6", "√36 = 6"), ("√25 = 5", "√25 = 5"), ("√25 · √2 = 5√2", "(3 + 2)√2 = 5√2"),
    ("√9 · √3 = 3√3", "2√3 + √3 = 3√3"), ("0,4 · 5 = 2", "√4 = 2"), ("5 + 3 = 8", "√64 = 8"),
    ("4√3 − 2√3 = 2√3", "√4 · √3 = 2√3"), ("3. binom. Formel: 5 − 1 = 4", "3. binom. Formel: 7 − 3 = 4"),
    ("1. binom. Formel: 1 + 2√2 + 2", "2 + 2√2 + 1"), ("√9 · √2 · √x² = 3x√2", "3 · x · √2 = 3x√2"),
]
STERNE_PC = [1, 1, 1, 1, 1, 2, 2, 2, 3, 3]
neues_blatt("Partnercheck: Rechnen mit Wurzeln", "zu zweit – verschiedene Aufgaben, gleiche Ergebnisse")
absatz(d, "**So geht’s:** Faltet das Blatt in der Mitte. A rechnet nur Spalte A, B nur Spalte B. "
          "**In jeder Zeile muss bei euch beiden dasselbe Ergebnis herauskommen.** Vergleicht nach jeder Zeile. "
          "Stimmen eure Ergebnisse nicht überein, sucht gemeinsam den Fehler – erst dann geht es weiter. "
          "Gib Ergebnisse **exakt** an (z. B. 5√2, nicht 7,07).")
t = tabelle(d, ["", "Partner*in A", "Partner*in B"],
            [[f"{i + 1}  {sterne_text(s)}", a + "\n\n", b + "\n\n"] for i, ((a, b, _), s) in enumerate(zip(PC, STERNE_PC))],
            breiten_cm=[2.2, 7.6, 7.6])
absatz(d, "Satzbausteine bei Abweichung: „Bei mir kommt … heraus, weil …“ · „Ich glaube, der Fehler liegt bei …“ · "
          "„Wir können prüfen, indem wir … quadrieren / ausmultiplizieren.“", groesse=10)

neues_blatt("Lösung: Partnercheck")
tabelle(d, ["", "Rechenweg A", "Rechenweg B", "Ergebnis"],
        [[str(i + 1), wa, wb, e] for i, ((_, _, e), (wa, wb)) in enumerate(zip(PC, PC_WEG))],
        breiten_cm=[0.8, 6.6, 6.6, 3.4], schrift=10)

# ===================================================================== 4 Fehlerdetektiv
CHAT = [
    (1, "@nightowl09", "√(−25) = −5"),
    (1, "@goalkeeper7", "x² = 49  →  x = 7. Fertig!"),
    (1, "@lofi_beats", "√0,4 = 0,2"),
    (1, "@pixelqueen", "√49 ist irrational, weil da eine Wurzel steht."),
    (2, "@hakuna_matheta", "√(16 + 9) = √16 + √9 = 4 + 3 = 7"),
    (2, "@skater.boi", "√8 + √2 = √10"),
    (2, "@cloud_surfer", "Mein TR zeigt √2 = 1,414213562. Die Zahl bricht ab, also ist √2 rational."),
    (3, "@wurzelzwerg", "(√3 + 1)² = 3 + 1 = 4"),
]
CHAT_LOES = [
    "Keine Zahl ergibt quadriert −25 (auch (−5)² = 25). √(−25) ist nicht definiert.",
    "Auch (−7)² = 49. Lösungsmenge L = {−7; 7}.",
    "0,2² = 0,04, nicht 0,4. Richtig wäre √0,04 = 0,2; √0,4 ≈ 0,63 ist irrational.",
    "√49 = 7 ist eine natürliche (also auch rationale) Zahl. Nur Wurzeln aus Nicht-Quadratzahlen sind irrational.",
    "Wurzeln darf man bei Summen nicht einzeln ziehen: √(16 + 9) = √25 = 5.",
    "Summen-Fehler: √8 = 2√2, also √8 + √2 = 2√2 + √2 = 3√2 (≈ 4,24; √10 ≈ 3,16).",
    "Der TR rundet. 1,414213562² endet auf der Ziffer 4 (2 · 2 = 4), ist also nicht genau 2. √2 ist irrational.",
    "Binomische Formel vergessen: (√3 + 1)² = 3 + 2√3 + 1 = 4 + 2√3.",
]
neues_blatt("Fehlerdetektiv: Der Klassenchat am Abend vor der Arbeit", "Kann man den Lösungen trauen?")
absatz(d, "**Auftrag:** (1) **Finde** in jeder Nachricht den Fehler. (2) **Erkläre** ihn mit einem Satzbaustein. "
          "(3) **Schreibe** die richtige Lösung auf.")
satzbausteine([
    "„Der Fehler liegt bei …“ · „Richtig wäre …, weil …“ · „Das kann man überprüfen, indem man … quadriert.“",
    "„Man darf … nicht …, denn …“ · „Ein Gegenbeispiel ist …“",
])
for nr, (s, wer, txt) in enumerate(CHAT, start=1):
    tb = d.add_table(rows=1, cols=2); _rahmen(tb, staerke=6)
    _nicht_trennen(tb.rows[0])
    c0, c1 = tb.cell(0, 0), tb.cell(0, 1)
    _schattierung(c0, GRAU)
    _text(c0.paragraphs[0], f"{nr}  {sterne_text(s)}", fett=True)
    absatz(c0, f"**{wer}:**  {txt}")
    _text(c1.paragraphs[0], "Fehler und richtige Lösung:")
    absatz(c1, ""); absatz(c1, "")
    c0.width, c1.width = Cm(8.2), Cm(9.2)
    sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = Pt(6)

neues_blatt("Lösung: Fehlerdetektiv")
tabelle(d, ["Nr.", "Nachricht", "Fehler und richtige Lösung"],
        [[str(i + 1), c[2], l] for i, (c, l) in enumerate(zip(CHAT, CHAT_LOES))],
        breiten_cm=[1, 6.4, 10], schrift=10)

# ===================================================================== 5 Expertenblätter A–D
def expertenblatt(kennung, thema, kontext, info, aufgaben, erklaer, kontroll, wortspeicher=None):
    neues_blatt(f"Expertenblatt {kennung}: {thema}",
                f"Gruppenpuzzle „Wurzel-Expert*innen“ · Lerntheke Pflichtstation P{'ABCD'.index(kennung) + 1}")
    absatz(d, kontext)
    merkkasten(d, "Das musst du wissen", info)
    for nr, (s, text, teile, spalten) in enumerate(aufgaben, start=1):
        aufgabe(d, nr, text, sterne=s, teile=teile, spalten=spalten)
    if wortspeicher:
        satzbausteine(wortspeicher)
    merkkasten(d, "Für deine Stammgruppe", [
        "**Erklärauftrag:** " + erklaer,
        "Schreibe dir einen **Erklärzettel** (höchstens 5 Stichpunkte + 1 Beispiel).",
        "**Kontrollfrage** für deine Gruppe: " + kontroll,
    ])


expertenblatt(
    "A", "Quadratwurzeln und Gleichungen x² = a",
    "**Problem:** Ein quadratisches Pixel-Art-Sprite besteht aus 256 Pixeln. Wie viele Pixel ist es breit? "
    "Und welche Zahl ergibt quadriert 256 – gibt es nur eine?",
    ["Für a ≥ 0 ist **√a** die **nicht negative** Zahl, die quadriert a ergibt. Beispiel: √256 = 16, denn 16² = 256.",
     "Aus negativen Zahlen kann man keine Quadratwurzel ziehen: √(−4) gibt es nicht.",
     "Die Gleichung **x² = a** hat für a > 0 **zwei** Lösungen: x = √a und x = −√a. "
     "Beispiel: x² = 256 → L = {−16; 16}. Für a = 0: L = {0}. Für a < 0: L = { }."],
    [(1, "Berechne ohne Taschenrechner.", ["√196", "√0,49", "√(81/100)", "√1,44"], 4),
     (1, "Bestimme die Lösungsmenge.", ["x² = 64", "x² = 0", "x² = −9"], 3),
     (2, "Ein quadratischer Instagram-Post hat 1 166 400 Pixel. Berechne ohne Taschenrechner, wie viele Pixel er breit ist. "
         "Tipp: 1 166 400 = 11 664 · 100 und 108² = 11 664.", None, 1),
     (2, "Bestimme die Lösungsmenge.", ["x² − 2,25 = 0", "2x² = 50"], 2),
     (3, "Jemand behauptet: „√(a²) = a gilt für jede Zahl a.“ Prüfe die Aussage mit a = −3 und gib an, für welche "
         "Zahlen a sie stimmt.", None, 1)],
    "Warum hat x² = 25 zwei Lösungen, obwohl √25 nur einen Wert hat?",
    "Löse x² = 1,69.",
)

expertenblatt(
    "B", "Wurzeln näherungsweise bestimmen",
    "**Problem:** Wie lange fällt man vom Sprungturm im Freibad? Ohne Luftwiderstand gilt näherungsweise "
    "**t ≈ √(h : 5)** (t in Sekunden, h in Metern; dabei wurde die Fallbeschleunigung auf 10 m/s² gerundet). "
    "Vom 7,5-m-Turm fällt man also etwa √1,5 Sekunden – aber wie viel ist das?",
    ["**Intervallschachtelung:** Man sucht Zahlen, deren Quadrate knapp unter und knapp über dem Radikanden liegen.",
     "Beispiel √2:  1² = 1 < 2 < 4 = 2²  →  1 < √2 < 2;   1,4² = 1,96 < 2 < 2,25 = 1,5²  →  1,4 < √2 < 1,5;   "
     "1,41² = 1,9881 < 2 < 2,0164 = 1,42²  →  1,41 < √2 < 1,42.",
     "Mit jedem Schritt kennt man eine Nachkommastelle mehr. Fertig ist man nie: √2 hat unendlich viele Nachkommastellen."],
    [(1, "Gib an, zwischen welchen natürlichen Zahlen die Wurzel liegt.", ["√40", "√90", "√150"], 3),
     (2, "Bestimme √1,5 mit Intervallschachtelung auf zwei Nachkommastellen. Wie lange fällt man also vom 7,5-m-Turm?", None, 1),
     (2, "Berechne mit der Formel die Fallzeit vom 10-m-Turm und vom 3-m-Brett (Taschenrechner erlaubt, auf "
         "Hundertstel runden).", None, 1),
     (3, "Der 10-m-Turm ist mehr als dreimal so hoch wie das 3-m-Brett. **Begründe**, warum die Fallzeit trotzdem "
         "nicht dreimal so lang ist. Wie hoch müsste ein Turm sein, damit man doppelt so lange fällt wie vom 10-m-Turm?",
      None, 1)],
    "Wie funktioniert die Intervallschachtelung? Zeige es an √1,5.",
    "Zwischen welchen Zehnteln liegt √3?",
)

expertenblatt(
    "C", "Rationale und irrationale Zahlen",
    "**Problem:** Der Taschenrechner zeigt √2 = 1,414213562. Ist das der genaue Wert – und kann man √2 als Bruch schreiben?",
    ["**Rationale Zahlen** (ℚ) lassen sich als Bruch schreiben. Als Dezimalzahl sind sie **abbrechend** "
     "(0,75 = 3/4) oder **periodisch** (0,3̅ = 1/3).",
     "**Irrationale Zahlen** sind unendlich lang und **nicht periodisch**, z. B. √2, √3, π. "
     "√n ist für natürliche Zahlen n irrational, wenn n keine Quadratzahl ist.",
     "**Endziffern-Trick:** 1,414213562 endet auf 2, ihr Quadrat endet deshalb auf 4 (2 · 2 = 4) – es kann nicht genau 2 sein.",
     "Zahlbereiche:  ℕ ⊂ ℤ ⊂ ℚ ⊂ ℝ.  Rationale und irrationale Zahlen bilden zusammen die reellen Zahlen ℝ.",
     "**Satzbausteine:** „… ist rational, weil …“ · „… ist irrational, weil …“ · „Ein Gegenbeispiel ist …, denn …“"],
    [(1, "Entscheide: rational oder irrational?", ["√36", "√37", "0,75", "0,3̅", "π", "√(9/4)", "2,1010010001…"], 4),
     (1, "Schreibe als vollständig gekürzten Bruch:  0,75;  0,3̅;  √(9/4).", None, 1),
     (2, "Ein DIN-A4-Blatt ist 297 mm lang und 210 mm breit. **Berechne** das Verhältnis 297 : 210 und **entscheide**, "
         "ob es genau √2 ist.", None, 1),
     (2, "Der Taschenrechner zeigt √3 = 1,7320508. **Zeige** mit dem Endziffern-Trick, dass das nicht der genaue Wert ist.",
      None, 1),
     (3, "Stimmt das? **Begründe** oder **widerlege** mit einem Gegenbeispiel.",
      ["Die Summe zweier irrationaler Zahlen ist immer irrational.",
       "Das Produkt zweier irrationaler Zahlen ist immer irrational.",
       "Jede Wurzel ist irrational."], 1)],
    "Woran erkennt man, ob eine Zahl rational oder irrational ist?",
    "Ordne in den kleinsten Zahlbereich ein:  −4;  2/3;  √5;  √16.",
)

expertenblatt(
    "D", "Rechnen mit Wurzeln (Wurzelgesetze)",
    "**Problem:** In einem „Mathe-Hack“-Video heißt es: „√(16 + 9) = √16 + √9 = 7 – so rechnest du doppelt so schnell!“ "
    "Stimmt das?",
    ["Für a, b ≥ 0:  **√a · √b = √(a · b)**   und (für b > 0)  **√a : √b = √(a : b)**.",
     "**Teilweise Wurzel ziehen:** Radikand so zerlegen, dass ein Faktor eine Quadratzahl ist:  √72 = √36 · √2 = 6√2.",
     "**Zusammenfassen** geht nur bei gleichen Wurzeln:  3√2 + 5√2 = 8√2.",
     "**Achtung:** Für Summen gibt es kein solches Gesetz:  √(a + b) ≠ √a + √b  (außer wenn a = 0 oder b = 0)."],
    [(1, "Berechne.", ["√2 · √32", "√3 · √27", "√98 : √2", "√0,2 · √20"], 4),
     (1, "Ziehe teilweise die Wurzel.", ["√20", "√45", "√200"], 3),
     (2, "Fasse zusammen.", ["√18 + √50", "√12 + √27 − √3"], 2),
     (2, "Prüfe den „Mathe-Hack“ aus dem Video. Finde ein Zahlenpaar, für das √(a + b) = √a + √b doch stimmt.", None, 1),
     (3, "Ein quadratisches Beet im Schulgarten hat 50 m². Ein zweites quadratisches Beet soll doppelt so viel Fläche haben. "
         "**Berechne** beide Seitenlängen exakt und **erkläre**, mit welchem Faktor sich die Seitenlänge ändert, "
         "wenn sich die Fläche verdoppelt.", None, 1)],
    "Welche Rechengesetze gelten für Wurzeln – und wo lauert die Falle?",
    "Sind √8 · √2 und √8 + √2 gleich groß?",
)

neues_blatt("Abschluss-Challenge für die Stammgruppe", "Gruppenpuzzle „Wurzel-Expert*innen“")
absatz(d, "Eure Klasse will für das Schulfest ein **quadratisches Banner mit 2 m² Fläche** drucken lassen. "
          "Jede*r Expert*in übernimmt die Teilaufgabe zu seinem Thema und erklärt sie den anderen.")
aufgabe(d, 1, "**(A)** Gib die Seitenlänge des Banners exakt an.", sterne=1, platz=1)
aufgabe(d, 2, "**(C)** Ist die Seitenlänge eine rationale Zahl? Begründe.", sterne=2, platz=2)
aufgabe(d, 3, "**(B)** Die Druckerei braucht die Seitenlänge in cm. Bestimme sie mit Intervallschachtelung auf ganze cm "
              "(abgerundet) und erkläre, warum das Banner dann etwas kleiner als 2 m² ist.", sterne=2, platz=3)
aufgabe(d, 4, "**(D)** Ein zweites Banner soll 8 m² groß sein. Gib seine Seitenlänge vereinfacht an. "
              "Wie viel Mal so lang ist sie wie beim ersten Banner?", sterne=3, platz=2)
absatz(d, "**Zurück zur Leitfrage:** Formuliert gemeinsam einen Antwortsatz.")
schreibplatz(d, 2)

# ---------------------------------------------------------------- Lösungen A–D + Challenge
neues_blatt("Lösungen: Expertenblätter A–D und Challenge")
ueberschrift(d, "A  Quadratwurzeln und Gleichungen", 3)
liste(d, [
    "1  a) 14   b) 0,7   c) 9/10   d) 1,2",
    "2  a) L = {−8; 8}   b) L = {0}   c) L = { }  (kein Quadrat ist negativ)",
    "3  √1 166 400 = √11 664 · √100 = 108 · 10 = **1080 Pixel**",
    "4  a) x² = 2,25 → L = {−1,5; 1,5}   b) x² = 25 → L = {−5; 5}",
    "5  a = −3: √((−3)²) = √9 = 3 ≠ −3. Die Aussage gilt nur für **a ≥ 0** (allgemein: √(a²) = |a|, der Betrag von a).",
    "Kontrollfrage: L = {−1,3; 1,3}",
])
ueberschrift(d, "B  Wurzeln näherungsweise bestimmen", 3)
liste(d, [
    "1  a) 6 < √40 < 7   b) 9 < √90 < 10   c) 12 < √150 < 13",
    "2  1,2² = 1,44 < 1,5 < 1,69 = 1,3²;  1,22² = 1,4884 < 1,5 < 1,5129 = 1,23²  →  1,22 < √1,5 < 1,23. "
    "Fallzeit vom 7,5-m-Turm **etwa 1,22 s**.",
    "3  10 m: t ≈ √2 ≈ **1,41 s**;  3 m: t ≈ √0,6 ≈ **0,77 s**.",
    "4  Die Höhe steht unter der Wurzel: Vervierfacht man h, verdoppelt sich t, denn √(4h : 5) = 2 · √(h : 5). "
    "Dreifache Höhe gibt nur etwa √3 ≈ 1,7-fache Zeit. Doppelte Fallzeit: **40 m** hoher Turm.",
    "Kontrollfrage: 1,7² = 2,89 < 3 < 3,24 = 1,8²  →  1,7 < √3 < 1,8",
])
ueberschrift(d, "C  Rationale und irrationale Zahlen", 3)
liste(d, [
    "1  rational: √36 = 6, 0,75, 0,3̅, √(9/4) = 3/2;  irrational: √37, π, 2,1010010001… (nicht periodisch)",
    "2  0,75 = 3/4;  0,3̅ = 1/3;  √(9/4) = 3/2",
    "3  297 : 210 = 99/70 = 1,41428…; √2 = 1,41421… Das Verhältnis ist ein Bruch, also rational – √2 ist irrational. "
    "Es ist also **nicht genau √2**, sondern ein auf ganze mm gerundeter Wert (sehr nah dran).",
    "4  1,7320508 endet auf 8; 8 · 8 = 64, also endet das Quadrat auf 4 und kann nicht genau 3 sein.",
    "5  a) falsch: √2 + (−√2) = 0   b) falsch: √2 · √8 = √16 = 4   c) falsch: √49 = 7",
    "Kontrollfrage: −4 ∈ ℤ;  2/3 ∈ ℚ;  √5 irrational (ℝ);  √16 = 4 ∈ ℕ",
])
ueberschrift(d, "D  Rechnen mit Wurzeln", 3)
liste(d, [
    "1  a) √64 = 8   b) √81 = 9   c) √49 = 7   d) √4 = 2",
    "2  a) 2√5   b) 3√5   c) 10√2",
    "3  a) 3√2 + 5√2 = 8√2   b) 2√3 + 3√3 − √3 = 4√3",
    "4  √(16 + 9) = √25 = 5, nicht 7 – der Hack ist falsch. Es stimmt nur, wenn a = 0 oder b = 0, z. B. √(0 + 9) = 0 + 3.",
    "5  √50 = 5√2 m ≈ 7,07 m;  √100 = 10 m. Faktor 10 : 5√2 = √2 ≈ 1,41: "
    "Doppelte Fläche bedeutet √2-fache Seitenlänge (wie bei DIN A4 → DIN A3).",
    "Kontrollfrage: √8 · √2 = 4,  √8 + √2 = 3√2 ≈ 4,24 – nicht gleich.",
])
ueberschrift(d, "Abschluss-Challenge", 3)
liste(d, [
    "1  Seitenlänge √2 m",
    "2  Nein, √2 ist irrational (2 ist keine Quadratzahl; unendlicher, nicht periodischer Dezimalbruch).",
    "3  1,41² = 1,9881 < 2 < 2,0164 = 1,42²  →  **141 cm**. Abgerundet: 1,41 m · 1,41 m = 1,9881 m² < 2 m².",
    "4  √8 = 2√2 m ≈ 2,83 m; 2√2 : √2 = **2-mal so lang** (vierfache Fläche → doppelte Seite).",
])

# ===================================================================== 6 Laufzettel
neues_blatt("Laufzettel Lerntheke: Wurzeln und reelle Zahlen", "Was muss ich bis zur Klassenarbeit noch üben?")
absatz(d, "**1. Schätze dich ein** – vorher (Beginn) und nachher (Ende der Stunde).  "
          "++ = sicher · + = ziemlich sicher · – = unsicher")
tabelle(d, ["Ich kann …", "vorher\n++   +   –", "nachher\n++   +   –", "Station"], [
    ["Quadratwurzeln ohne Taschenrechner bestimmen und Gleichungen x² = a lösen.", "☐  ☐  ☐", "☐  ☐  ☐", "P1 (A), W1"],
    ["Wurzeln mit Intervallschachtelung näherungsweise bestimmen.", "☐  ☐  ☐", "☐  ☐  ☐", "P2 (B)"],
    ["rationale und irrationale Zahlen unterscheiden und Zahlbereichen zuordnen.", "☐  ☐  ☐", "☐  ☐  ☐", "P3 (C), W2"],
    ["mit den Wurzelgesetzen rechnen und teilweise die Wurzel ziehen.", "☐  ☐  ☐", "☐  ☐  ☐", "P4 (D), W3, W5"],
    ["typische Fehler erkennen und erklären.", "☐  ☐  ☐", "☐  ☐  ☐", "W4"],
], breiten_cm=[8.6, 2.9, 2.9, 3.0], schrift=10)
absatz(d, "**2. Bearbeite die Stationen.** Pflicht: von jeder P-Station mindestens die ★-Aufgaben. "
          "Beginne mit der Station, bei der du „–“ angekreuzt hast. Kontrolliere an der Lösungsstation.")
tabelle(d, ["Station", "Thema / Methode", "Sozialform", "erledigt", "kontrolliert"], [
    ["P1", "Quadratwurzeln und x² = a (Expertenblatt A)", "allein", "☐", "☐"],
    ["P2", "Wurzeln näherungsweise bestimmen (Expertenblatt B)", "allein", "☐", "☐"],
    ["P3", "Rationale und irrationale Zahlen (Expertenblatt C)", "allein", "☐", "☐"],
    ["P4", "Rechnen mit Wurzeln (Expertenblatt D)", "allein", "☐", "☐"],
    ["W1", "Wurzel-Domino", "zu zweit / dritt", "☐", "☐"],
    ["W2", "Wo wohnt die Zahl? (Sortierspiel)", "zu dritt / viert", "☐", "☐"],
    ["W3", "Partnercheck: Rechnen mit Wurzeln", "zu zweit", "☐", "☐"],
    ["W4", "Fehlerdetektiv: Klassenchat", "allein → zu zweit", "☐", "☐"],
    ["W5", "Partnerspiel: Wer findet schneller ein Wurzelpaar?", "zu zweit", "☐", "☐"],
], breiten_cm=[1.6, 8.4, 3.4, 2, 2], schrift=10)
merkkasten(d, "3. Exit-Ticket (allein, ohne Taschenrechner)", [
    "a) Ziehe teilweise die Wurzel: √75 = ____________",
    "b) Ist √0,36 rational oder irrational? Begründe. ______________________________________",
    "c) Diese Station mache ich vor der Klassenarbeit noch einmal: ____________",
])
absatz(d, "_Lösung Exit-Ticket (für die Lehrkraft): a) √25 · √3 = 5√3   b) rational, denn √0,36 = 0,6 = 3/5_", groesse=9)

d.save("reihen/01-wurzeln/methodenpool-reelle-zahlen.docx")
print("ok")
